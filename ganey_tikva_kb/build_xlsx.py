"""Build the knowledge-items Excel from items.json.
items.json: list of {title, category, source_url, updated, body:[blocks]}
block: {"h": text} heading | {"p": text} paragraph (supports **bold**) | {"li": [..]} bullets
       | {"img": url, "alt": text} | {"file": url, "text": name} | {"link": url, "text": name}
Usage: python3 -I build_xlsx.py items.json out.xlsx"""
import sys, json, re
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

B = InlineFont(b=True, rFont="Arial", sz=11); N = InlineFont(rFont="Arial", sz=11)
H = InlineFont(b=True, rFont="Arial", sz=12, color="1F4E79"); L = InlineFont(rFont="Arial", sz=11, color="0563C1", u="single")

def rich(blocks):
    parts = []
    def add(t, f): parts.append(TextBlock(f, t))
    def inline(t):
        for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", t)):
            if seg: add(seg, B if i % 2 else N)
    for b in blocks:
        if parts: add("\n", N)
        if "h" in b: add(b["h"], H)
        elif "p" in b: inline(b["p"])
        elif "li" in b:
            for j, x in enumerate(b["li"]):
                if j: add("\n", N)
                add("• ", N); inline(x)
        elif "img" in b: add(f"🖼 תמונה: {b.get('alt','')} – {b['img']}", L)
        elif "file" in b: add(f"📎 {b.get('text','קובץ')}: {b['file']}", L)
        elif "link" in b: add(f"🔗 {b.get('text','קישור')}: {b['link']}", L)
    return CellRichText(parts)

def plain(blocks):
    return "\n".join(b.get("h") or b.get("p") or "\n".join(b.get("li", [])) or b.get("img") or b.get("file") or b.get("link") or "" for b in blocks)

items = json.load(open(sys.argv[1], encoding="utf-8"))
wb = Workbook(); ws = wb.active; ws.title = "פריטי מידע"; ws.sheet_view.rightToLeft = True
cols = [("קוד פריט", 10), ("שם פריט", 38), ("טקסט", 95), ("קטגוריה", 20), ("כתובת מקור", 40), ("קבצים מצורפים", 40), ("תאריך עדכון באתר", 14)]
ws.append([c for c, _ in cols])
for i, (_, w) in enumerate(cols, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
thin = Side(style="thin", color="BFBFBF")
for n, it in enumerate(items, 1):
    files = "\n".join(f"{b.get('text','')}: {b['file']}" for b in it["body"] if "file" in b)
    ws.append([n, it["title"], None, it.get("category", ""), it.get("source_url", ""), files, it.get("updated", "")])
    ws.cell(n + 1, 3).value = rich(it["body"])
    lines = plain(it["body"]).count("\n") + len(plain(it["body"])) // 110 + 1
    ws.row_dimensions[n + 1].height = min(409, 15 * lines + 6)
    if it.get("source_url"): ws.cell(n + 1, 5).hyperlink = it["source_url"]; ws.cell(n + 1, 5).style = "Hyperlink"
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if c.row > 1 and not isinstance(c.value, CellRichText) and c.style != "Hyperlink": c.font = Font(name="Arial", size=11)
for c in ws[1]:
    c.font = Font(name="Arial", bold=True, color="FFFFFF", size=12); c.fill = PatternFill("solid", fgColor="1F4E79")
    c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
wb.save(sys.argv[2]); print("saved", len(items), "items ->", sys.argv[2])
